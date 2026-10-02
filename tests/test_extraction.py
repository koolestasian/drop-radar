import tempfile
import unittest
from datetime import date, datetime, timezone
from pathlib import Path
from unittest.mock import patch

from radar.legacy import opportunity_monitor as monitor
from openpyxl import load_workbook

FIXTURE = Path(__file__).parent / "fixtures" / "sample_items.json"


def record(**values):
    row = {header: "" for header in monitor.HEADERS}
    row.update({"Actioned?": "No", "Source Type": "Story"})
    row.update(values)
    return row


class OrganizationTests(unittest.TestCase):
    def org(self, url):
        return monitor.organization_from_url(url)

    def test_ats_hosts_name_the_company_not_the_ats(self):
        self.assertEqual(self.org("https://jobs.ashbyhq.com/ramp/acf6b28d-767f-483f-8ff2-114620"), "Ramp")
        self.assertEqual(self.org("https://jobs.smartrecruiters.com/Wise/744000151030889-graduate"), "Wise")
        self.assertEqual(self.org("https://jobs.smartrecruiters.com/LinkedIn3/744000151447279-s"), "LinkedIn")
        self.assertEqual(
            self.org("https://job-boards.greenhouse.io/embed/job_app?for=doordashusa&token=7848317"),
            "DoorDash",
        )
        self.assertEqual(self.org("https://app.careerpuck.com/job-board/lyft/job/8817900002"), "Lyft")

    def test_slugs_map_to_display_names(self):
        self.assertEqual(self.org("https://capitalone.wd12.myworkdayjobs.com/en-US/Capital_One/job/x"), "Capital One")
        self.assertEqual(self.org("https://job-boards.greenhouse.io/scaleai/jobs/4736426005"), "Scale AI")
        self.assertEqual(self.org("https://careers.withwaymo.com/jobs/2027-summer-intern"), "Waymo")
        self.assertEqual(self.org("https://www.lifeatspotify.com/jobs/recsys-2026-intern"), "Spotify")
        self.assertEqual(self.org("https://uscareers-lennox.icims.com/jobs/54888/job"), "Lennox")
        self.assertEqual(self.org("https://www.tesla.com/careers/search/job/285089"), "Tesla")

    def test_generic_hosts_do_not_become_organizations(self):
        for url in (
            "https://tinyurl.com/easyf100-2026",
            "https://luma.com/ptsnzkkl",
            "https://docs.google.com/forms/d/e/abc/viewform",
            "https://hdnn.fa.us6.oraclecloud.com/hcmUI/CandidateExperience/en/sites/CX_1/job/1",
        ):
            self.assertEqual(self.org(url), "", url)

    def test_oracle_site_name_is_used(self):
        url = "https://icbpjb.fa.ocs.oraclecloud.com/hcmUI/CandidateExperience/en/sites/LazardStudentCareers/job/6606"
        self.assertEqual(self.org(url), "Lazard")

    def test_truncated_ocr_url_is_not_a_link(self):
        self.assertEqual(monitor.external_links(["https://ejko.fa.us2"]), [])

    def test_text_match_is_case_sensitive_for_common_words(self):
        self.assertEqual(monitor.organization_from_text("We offer visa sponsorship"), "")
        self.assertEqual(monitor.organization_from_text("Target audience: students"), "Target")
        self.assertEqual(monitor.organization_from_text("NVIDIA and Google are hiring"), "NVIDIA")


class TitleTests(unittest.TestCase):
    def test_title_from_url_slug(self):
        self.assertEqual(
            monitor.title_from_url(
                "https://www.amazon.jobs/en/jobs/10565667/software-development-engineer-intern-aws-database-2027-us",
                "Amazon",
            ),
            "Software Development Engineer Intern AWS Database 2027 US",
        )
        self.assertEqual(
            monitor.title_from_url(
                "https://moog.wd5.myworkdayjobs.com/en-US/MOOG_External_Career_Site/job/Intern--Software-Engineering_R-26-19948",
                "Moog",
            ),
            "Intern Software Engineering",
        )
        self.assertEqual(
            monitor.title_from_url(
                "https://www.google.com/about/careers/applications/jobs/results/"
                "107272685397910214-associate-product-manager-university-graduate-2027-start",
                "Google",
            ),
            "Associate Product Manager University Graduate 2027 Start",
        )

    def test_org_name_is_not_repeated_in_title(self):
        title = monitor.title_from_url(
            "https://mycareer.verizon.com/jobs/r-1101384/verizon-network-and-technology-data-science-summer-2027-internship",
            "Verizon",
        )
        self.assertEqual(title, "Network and Technology Data Science Summer 2027 Internship")

    def test_id_only_urls_have_no_title(self):
        self.assertEqual(monitor.title_from_url("https://www.tesla.com/careers/search/job/285089"), "")
        self.assertEqual(monitor.title_from_url("https://job-boards.greenhouse.io/figma/jobs/6200626004"), "")
        self.assertEqual(
            monitor.title_from_url("https://jobs.ashbyhq.com/ramp/acf6b28d-767f-483f-8ff2-114620b3c4d5"), ""
        )


    def test_ocr_sentence_fragments_are_not_titles(self):
        for line in (
            "career fair but we are definitely",
            "the world, we are looking for a Data Science Intern",
            "Data Engineer Intern and help build",
            "Thank you for applying to the Code for Good Hackathon",
        ):
            self.assertFalse(monitor.usable_text_title(line), line)
        self.assertTrue(monitor.usable_text_title("Lyft Content Systems Intern (Summer 2027)"))


class DeadlineTests(unittest.TestCase):
    reference = date(2026, 9, 20)

    def deadline(self, text):
        return monitor.extract_deadline(text, self.reference)

    def test_common_phrasings(self):
        self.assertEqual(self.deadline("Application closes September 25, 2026"), "2026-09-25")
        self.assertEqual(self.deadline("Application closes October 9"), "2026-10-09")
        self.assertEqual(self.deadline("Apply Before 12/30/2026, 04:00 PM"), "2026-12-30")
        self.assertEqual(self.deadline("Registration closes on October 2, 2026 at noon"), "2026-10-02")
        self.assertEqual(self.deadline("Registration Deadline: Oct 2, 2026 12:00 PM"), "2026-10-02")
        self.assertEqual(self.deadline("Apply by Oct. 3rd!"), "2026-10-03")
        self.assertEqual(
            self.deadline("Apply before 11:59 PM ET on Thursday, October 15"), "2026-10-15"
        )

    def test_no_date_no_deadline(self):
        self.assertEqual(self.deadline("Tell your friends before this fills up"), "")
        self.assertEqual(self.deadline("apply ASAP as job postings may close"), "")

    def test_yearless_date_early_in_year_rolls_forward(self):
        self.assertEqual(monitor.extract_deadline("Deadline: Jan 15", date(2026, 11, 1)), "2027-01-15")

    def test_year_before_the_post_is_treated_as_a_typo(self):
        self.assertEqual(self.deadline("submit before end of day on October 6, 2025"), "2026-10-06")


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        for name, value in (
            ("TRACKER_PATH", root / "tracker.xlsx"),
            ("LIVE_VIEW_PATH", root / "LATEST.md"),
        ):
            patcher = patch.object(monitor, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        monitor.create_workbook(monitor.TRACKER_PATH)


    def test_dashboard_holds_values_not_formulas(self):
        monitor.save_records([
            record(ID="one", Status="Open", Priority="High", **{"Actioned?": "Yes"}),
            record(ID="two", Status="Expired"),
        ])
        wb = load_workbook(monitor.TRACKER_PATH)
        values = {row[0].value: row[1].value for row in wb["Dashboard"].iter_rows(min_row=4)}
        wb.close()
        self.assertEqual(values["Total opportunities"], 2)
        self.assertEqual(values["Actioned"], 1)
        self.assertEqual(values["High priority"], 1)
        self.assertEqual(values["Expired / closed"], 1)

    def test_live_view_groups_rows(self):
        now = datetime(2026, 9, 30, 12, tzinfo=timezone.utc)
        monitor.write_live_view([
            record(ID="new", Opportunity="Fresh | Role", Priority="High", Status="Open",
                   **{"First Seen": "2026-09-29T10:00:00+00:00",
                      "Application / Registration Link": "https://example.com/jobs/1"}),
            record(ID="soon", Opportunity="Closing Role", Deadline="2026-10-03", Status="New",
                   **{"First Seen": "2026-09-01T10:00:00+00:00"}),
            record(ID="old", Opportunity="Old Role", Status="New",
                   **{"First Seen": "2026-08-01T10:00:00+00:00",
                      "Instagram Source": "https://www.instagram.com/stories/zero2sudo/"}),
            record(ID="done", Opportunity="Done Role", **{"Actioned?": "Yes",
                                                          "First Seen": "2026-09-29T10:00:00+00:00"}),
            record(ID="gone", Opportunity="Gone Role", Status="Expired",
                   **{"First Seen": "2026-09-01T10:00:00+00:00"}),
        ], now=now)
        text = monitor.LIVE_VIEW_PATH.read_text()
        sections = text.split("## ")
        self.assertIn("Closing Role", next(s for s in sections if s.startswith("⏰")))
        self.assertIn("Oct 3 (3d)", text)
        self.assertIn("Fresh \\| Role", next(s for s in sections if s.startswith("🆕")))
        self.assertIn("🔥", text)
        self.assertIn("Old Role", next(s for s in sections if s.startswith("📋")))
        self.assertIn("(expired)", text)
        self.assertIn("<summary><b>✅ Actioned (1)</b></summary>", text)
        self.assertIn("<summary><b>⌛ Past deadline or posting closed (1)</b></summary>", text)
        self.assertIn("[Apply ↗](<https://example.com/jobs/1>)", text)


if __name__ == "__main__":
    unittest.main()
