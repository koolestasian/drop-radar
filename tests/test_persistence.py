import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import opportunity_monitor as monitor
from openpyxl import load_workbook


class PersistenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.previous_cwd = os.getcwd()
        os.chdir(self.temp.name)
        self.addCleanup(self.temp.cleanup)
        self.addCleanup(os.chdir, self.previous_cwd)
        for name, value in (
            ("TRACKER_PATH", Path("tracker.xlsx")),
            ("STATUS_PATH", Path("monitor_status.json")),
            ("STATE_PATH", Path("monitor_state.json")),
            ("LIVE_VIEW_PATH", Path("LATEST.md")),
            ("APIFY_TOKEN", "test"),
            ("fetch_stories", lambda: []),
            ("fetch_posts", lambda: []),
            ("normalize_item", lambda item, source: item),
            ("migrate_workbook", lambda manual_fields=None: {"before": 0, "after": 0, "changed": False}),
            ("record_semantic_key", lambda row: row["ID"]),
            ("pull_google_manual_fields", lambda: {}),
            ("github_issue", lambda rows, batch_id="": True),
            ("JOB_PAGES_ENABLED", False),
            ("LLM_ENABLED", False),
            ("ENRICHMENT_PATH", Path("enrichment_cache.json")),
            ("ntfy_alert", lambda rows, batch_id="": True),
        ):
            patcher = patch.object(monitor, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def run_rows(self, rows, defer=True):
        with patch.object(monitor, "fetch_stories", return_value=rows):
            monitor.main(defer_notifications=defer)

    def test_repeat_then_append_preserves_existing_edits(self):
        self.run_rows([{"ID": "one", "Opportunity": "Internship"}])
        wb = load_workbook(monitor.TRACKER_PATH)
        wb["Opportunities"]["Q2"] = "Yes"
        wb["Opportunities"]["R2"] = "Keep my note"
        wb.save(monitor.TRACKER_PATH)
        wb.close()
        before = monitor.TRACKER_PATH.read_bytes()
        self.run_rows([{"ID": "one"}])
        self.assertEqual(before, monitor.TRACKER_PATH.read_bytes())
        self.assertEqual(json.loads(Path("monitor_status.json").read_text())["new_rows"], 0)
        self.run_rows([{"ID": "one"}, {"ID": "two"}, {"ID": "two"}])
        wb = load_workbook(monitor.TRACKER_PATH)
        self.assertEqual(wb["Opportunities"].max_row, 3)
        self.assertEqual(wb["Opportunities"]["Q2"].value, "Yes")
        self.assertEqual(wb["Opportunities"]["R2"].value, "Keep my note")
        wb.close()
        status = json.loads(Path("monitor_status.json").read_text())
        self.assertEqual((status["previous_rows"], status["new_rows"], status["total_rows"]), (1, 1, 2))

    def test_alert_exception_does_not_lose_saved_rows(self):
        def failing_alert(rows, batch_id=""):
            raise TimeoutError()
        with patch.object(monitor, "github_issue", failing_alert), patch.object(monitor, "ntfy_alert", failing_alert):
            with self.assertRaises(RuntimeError):
                self.run_rows([{"ID": "one"}], defer=False)
        self.assertEqual(monitor.existing_ids(), {"one"})

    def test_deferred_alerts_not_sent_before_commit(self):
        with patch.object(monitor, "github_issue") as notify:
            self.run_rows([{"ID": "one"}])
            notify.assert_not_called()
        state = json.loads(Path("monitor_state.json").read_text())
        self.assertEqual(state["pending_batches"][0]["rows"][0]["ID"], "one")

    def test_scrape_failure_is_not_reported_as_success(self):
        with patch.object(monitor, "fetch_stories", side_effect=RuntimeError("failed")):
            with self.assertRaises(RuntimeError):
                monitor.main(defer_notifications=True)
        self.assertFalse(Path("monitor_status.json").exists())

    def test_saved_row_verification_detects_lost_write(self):
        with patch.object(monitor, "append_rows"):
            with self.assertRaisesRegex(RuntimeError, "Saved tracker IDs"):
                self.run_rows([{"ID": "one"}])

    def test_failed_alert_batch_is_retried_without_losing_state(self):
        row = {"ID": "one", "Opportunity": "Internship"}
        batch_id = monitor.queue_batch([row])
        attempts = []

        def flaky_issue(rows, current_batch_id=""):
            attempts.append(current_batch_id)
            if len(attempts) == 1:
                raise TimeoutError("temporary")
            return True

        with patch.object(monitor, "github_issue", flaky_issue):
            with self.assertRaisesRegex(RuntimeError, "remain pending"):
                monitor.deliver_pending_batches()
            monitor.deliver_pending_batches()

        state = json.loads(Path("monitor_state.json").read_text())
        self.assertEqual(attempts, [batch_id, batch_id])
        self.assertEqual(state["pending_batches"], [])
        self.assertIn(batch_id, state["delivered_batch_ids"])


if __name__ == "__main__":
    unittest.main()
